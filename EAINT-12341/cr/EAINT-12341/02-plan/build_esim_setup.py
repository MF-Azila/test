"""Build EAINT-12341_eSim_Setup.xlsx — what the tester types into eSim, per robot run.

    python build_esim_setup.py

The people (names / ICs) follow chainPeople() in
../03-spec/chain-scenarios.ts — keep the two in step.
Field names and response codes follow eAuto_eSimulator_Guides.xlsx (SSM_SIM_1 / SSM_SIM_2).
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "EAINT-12341_eSim_Setup.xlsx")
F = "Arial"
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = PatternFill("solid", fgColor="D9D9D9")
SEC = PatternFill("solid", fgColor="BFBFBF")
INPUT = PatternFill("solid", fgColor="FFFF00")
CALC = PatternFill("solid", fgColor="E2EFDA")
TOP = Alignment(wrap_text=True, vertical="top")
BACKUP_LINK = "https://cdn.eauto.my/FAQs_on_JPJeID_and_MyJPJ_App.pdf"

# Same order as the tester's table (02.10.2026).
RUNS = [
    dict(no=1, chain="CH1", type="ROB", label="Sole Prop / Partnership", tin="D", people=4, foreign=0, remark="Main test, run first"),
    dict(no=2, chain="CH2", type="ROB", label="Sole Prop / Partnership", tin="D", people=5, foreign=0, remark=""),
    dict(no=3, chain="CH6", type="ROB", label="Sole Prop / Partnership", tin="D (or IG)", people=4, foreign=0, remark="Stops at Submitted"),
    dict(no=4, chain="CH3", type="ROC", label="Sdn Bhd / Bhd", tin="C", people=3, foreign=1, remark="Has 1 foreign director"),
    dict(no=5, chain="CH8", type="ROC", label="Sdn Bhd / Bhd", tin="C", people=4, foreign=0, remark="Normal Sdn Bhd case"),
    dict(no=6, chain="CH4", type="LLP", label="LLP", tin="PT", people=4, foreign=0, remark="Check LLP0035174-LGN first"),
    dict(no=7, chain="CH5", type="TRADING", label="Business Trading (Sabah)", tin="TIN only", people=0, foreign=0, remark="No eSim. Fill ONLY the TIN cell."),
    dict(no=8, chain="CH7", type="TRADING", label="Business Trading (Sarawak)", tin="TIN only", people=0, foreign=0, remark="No eSim. Fill ONLY the TIN cell."),
    dict(no=9, chain="Spare 1", type="ROB", label="Sole Prop / Partnership", tin="D", people=0, foreign=0, remark="Only if a run fails"),
    dict(no=10, chain="Spare 2", type="ROB", label="Sole Prop / Partnership", tin="D", people=0, foreign=0, remark="Only if a run fails"),
    dict(no=11, chain="Spare 3", type="ROC", label="Sdn Bhd / Bhd", tin="C", people=0, foreign=0, remark="Only if a run fails"),
    dict(no=12, chain="Spare 4", type="LLP", label="LLP", tin="PT", people=0, foreign=0, remark="If available"),
]
SUFFIX = {"ROB": "ENTERPRISE", "ROC": "SDN BHD", "LLP": "PLT"}


def people(n, foreign):
    """Mirror of chainPeople(): Malaysians first, foreigners last."""
    out = []
    for i in range(n):
        f = i >= n - foreign
        ic = f"A1234567{i}" if f else f"{690501 + i}13{7631 + i}"
        dashed = ic if f else f"{ic[:6]}-{ic[6:8]}-{ic[8:]}"
        out.append(dict(name=f"QA DIRECTOR {chr(65 + i)}", plain=ic, dashed=dashed, foreign=f))
    return out


def c(ws, ref, v, bold=False, fill=None, border=True, italic=False, color=None, size=10):
    x = ws[ref]
    x.value = v
    x.font = Font(name=F, bold=bold, italic=italic, color=color, size=size)
    x.alignment = TOP
    if fill:
        x.fill = fill
    if border:
        x.border = BOX
    return x


wb = Workbook()

# ── How to use ───────────────────────────────────────────────────────────
ws = wb.active
ws.title = "How to use"
ws.column_dimensions["A"].width = 110
lines = [
    ("eSim setup for EAINT-12341 — what to type, run by run", True),
    ("", False),
    ("1. Open tab '1 Companies'. Fill ONLY the YELLOW cells: ROC/ROB number exactly as the checker sheet shows it (with the letter, e.g. 639691-H), the new 12-digit number, the TIN and the checker row.", False),
    ("   Only PASS rows. Do not use IP0581553-U, 003270488-X, 1267419-A (already used by Azfar).", False),
    ("2. The green cells work themselves out (number without the letter, the letter, the values for eSim).", False),
    ("3. For each run, open its tab (CH1, CH2, …). Log in to eSim: https://172.30.202.114:9089/esim/login (VPN on).", False),
    ("4. Follow the tab top to bottom: Part A = screen 'Ssm Enquiry Resp', Part B = 'Ssm Rob / Roc Submission Resp'.", False),
    ("   Edit an existing record (as the eSimulator guide says). Inside the big text boxes, change ONLY the tags listed — leave everything else as it is.", False),
    ("5. People: keep ONE existing person block, copy it until there are exactly as many blocks as the tab lists, then change the name and IC in each. Delete any extra blocks.", False),
    ("6. Save both screens, then tell QA support 'eSim ready for CH1' and start the robot (npm run ch1).", False),
    ("", False),
    ("Set up eSim for ONE run at a time, just before that run. Never reuse a company, even if a run fails — take a Spare row.", True),
    ("Send this file back to QA support after filling tab 1, so the same numbers go into the robot.", True),
    ("", False),
    ("Not confirmed yet (will be updated): the foreign director ID format (A1234567x used), and the LLP screen (not in the eSimulator guide).", False),
    ("Business Trading (CH5 Sabah, CH7 Sarawak) has no eSim setup — fill only its TIN in tab 1.", False),
]
for i, (t, b) in enumerate(lines, start=1):
    c(ws, f"A{i}", t, bold=b, border=False, size=12 if i == 1 else 10)

# ── 1 Companies ──────────────────────────────────────────────────────────
co = wb.create_sheet("1 Companies")
heads = ["No", "Robot run", "Company type", "TIN must start with", "People (robot expects)",
         "ROC / ROB No. (with letter)", "New 12-digit No.", "TIN", "Row in checker sheet", "Remarks",
         "eSim Roc (no letter)", "Check letter"]
widths = [5, 10, 22, 12, 26, 24, 18, 18, 12, 26, 20, 10]
for j, (h, w) in enumerate(zip(heads, widths), start=1):
    col = chr(64 + j)
    co.column_dimensions[col].width = w
    c(co, f"{col}3", h, bold=True, fill=CALC if j >= 11 else HDR)
c(co, "A1", "Fill the YELLOW cells only. Green = worked out for you.", bold=True, border=False)
row_of = {}
for i, r in enumerate(RUNS, start=4):
    row_of[r["chain"]] = i
    ppl = (f"{r['people'] - r['foreign']} Malaysian" + (f" + {r['foreign']} foreigner" if r["foreign"] else "")) if r["people"] else ("2 directors (robot types them)" if r["type"] == "TRADING" else "—")
    for j, v in enumerate([r["no"], r["chain"], r["label"], r["tin"], ppl], start=1):
        c(co, f"{chr(64 + j)}{i}", v)
    for col in "FGHI":
        c(co, f"{col}{i}", None, fill=INPUT)
    c(co, f"J{i}", r["remark"])
    # Strip a trailing "-<letter>" check letter; LLP "-LGN" stays whole.
    c(co, f"K{i}", f'=IF(F{i}="","",IF(AND(LEN(F{i})>2,MID(F{i},LEN(F{i})-1,1)="-",ISERROR(VALUE(RIGHT(F{i},1)))),LEFT(F{i},LEN(F{i})-2),F{i}))', fill=CALC)
    c(co, f"L{i}", f'=IF(OR(F{i}="",K{i}=F{i}),"",RIGHT(F{i},1))', fill=CALC)
ex = 4 + len(RUNS) + 1
for j, v in enumerate(["e.g.", "EXAMPLE", "Sdn Bhd / Bhd", "C", "—", "639691-H", "200401001188", "C11400448100", 21,
                       "EXAMPLE ONLY — format, do not use", "639691", "H"], start=1):
    c(co, f"{chr(64 + j)}{ex}", v, italic=True, color="808080")
co.freeze_panes = "A4"

# ── one tab per run ──────────────────────────────────────────────────────
def run_tab(r):
    ch = r["chain"]
    i = row_of[ch]
    t = r["type"]
    ws = wb.create_sheet(ch)
    for col, w in zip("ABC", (34, 62, 58)):
        ws.column_dimensions[col].width = w
    name = f"QA 12341 {ch} {SUFFIX[t]}"
    roc = f"'1 Companies'!K{i}"
    full = f"'1 Companies'!F{i}"
    new = f"'1 Companies'!G{i}"
    letter = f"'1 Companies'!L{i}"
    need = lambda ref: f'IF({ref}="","(fill tab 1 first)",{ref})'
    sub = {"ROB": "Ssm Rob Submission Resp", "ROC": "Ssm Roc Submission Resp", "LLP": "Ssm Llp Submission Resp (name to confirm)"}[t]
    c(ws, "A1", f"{ch} — {r['label']} — {r['remark'] or 'set up just before running ' + ch.lower()}", bold=True, border=False, size=12)
    c(ws, "A2", f"Company name the robot types: {name}", border=False)
    row = 4

    def sec(title):
        nonlocal row
        ws.merge_cells(f"A{row}:C{row}")
        c(ws, f"A{row}", title, bold=True, fill=SEC)
        row += 1
        for j, h in enumerate(["Field", "Type this", "Note"]):
            c(ws, f"{'ABC'[j]}{row}", h, bold=True, fill=HDR)
        row += 1

    def line(field, value, note="", calc=False):
        nonlocal row
        c(ws, f"A{row}", field)
        c(ws, f"B{row}", value, fill=CALC if calc else None)
        c(ws, f"C{row}", note)
        row += 1

    sec("Part A — screen 'Ssm Enquiry Resp' (edit a record, then Save)")
    line("Response Code", "004")
    line("Roc", f"={need(roc)}", "Number WITHOUT the letter", True)
    line("Check Digit (if the screen has it)", f'=IF({full}="","(fill tab 1 first)",IF({letter}="","(leave empty)",{letter}))', "The letter only. Without it eAuto shows '-A'.", True)
    line("New Roc", f"={need(new)}", "12 digits", True)
    line("Company Name", name, "Same as the robot types")
    line("Delay", "0")
    row += 1
    sec(f"Part B — screen '{sub}' (edit a record, then Save)")
    line("Response Code", "005")
    line("Roc", f"={need(roc)}", "Same as Part A", True)
    line("Check Digit (if the screen has it)", f'=IF({full}="","(fill tab 1 first)",IF({letter}="","(leave empty)",{letter}))', "Same as Part A", True)
    line("New Roc", f"={need(new)}", "Same as Part A", True)
    line("Company Name", name)
    line("Backup Link", BACKUP_LINK, "From the eSimulator guide")
    row += 1
    if t == "ROB":
        sec("Part B — text box 'Business Info': change ONLY these tags")
        line("<BUSINESS_REF_NO>", f'="<BUSINESS_REF_NO>"&{need(full)}&"</BUSINESS_REF_NO>"', "Number WITH the letter", True)
        line("<NEW_REF_NO>", f'="<NEW_REF_NO>"&{need(new)}&"</NEW_REF_NO>"', "", True)
        line("<ROB_STATUS>", "<ROB_STATUS>ACTIVE</ROB_STATUS>")
        row += 1
        sec(f"Part B — text box 'Current Owner Info': exactly {r['people']} owner blocks")
        icfield, key = "<IC_NEW> (no dashes)", "plain"
    elif t == "ROC":
        sec("Part B — text box 'Page 1': change ONLY these tags")
        line("<REF_NO>", f'="<REF_NO>"&{need(full)}&"</REF_NO>"', "Number WITH the letter", True)
        line("<NEW_REF_NO>", f'="<NEW_REF_NO>"&{need(new)}&"</NEW_REF_NO>"', "", True)
        line("<ROC_STATUS>", "<ROC_STATUS>EXISTING</ROC_STATUS>")
        row += 1
        sec(f"Part B — text box 'Page 3': exactly {r['people']} director blocks (designation DIRECTOR)")
        icfield, key = "<IC> (with dashes)", "dashed"
    else:
        sec("Part B — LLP is NOT in the eSimulator guide — send QA support a screenshot of one LLP record first")
        line("Registration number tag", f'={need(full)}', "Whatever tag holds the LLP number — WITH '-LGN'", True)
        line("New number tag", f"={need(new)}", "", True)
        line("Status tag", "active / existing", "Same idea as ROB / ROC — confirm on the screenshot")
        row += 1
        sec(f"Part B — text box 'Involvements': exactly {r['people']} partner blocks (type PT, no compliance officer)")
        icfield, key = "IC (no dashes)", "plain"
    for p in people(r["people"], r["foreign"]):
        line(f"Person {p['name'][-1]} — name", p["name"], "Foreign director — ID format to confirm" if p["foreign"] else "Malaysian")
        line(f"Person {p['name'][-1]} — {icfield}", p[key])
    row += 1
    c(ws, f"A{row}", "Then: Save, tell QA support 'eSim ready', and run the robot.", bold=True, border=False)


for r in RUNS:
    if r["people"]:
        run_tab(r)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("wrote", OUT)
